"""Build the LaundryTwin "Safe AI Executive Assistant" deck as a .pptx.

Everything is emitted as native PowerPoint shapes, text frames and tables
(no rasterised images), so individual elements can be selected and copied
straight into Canva.

Run:  .venv-pptx/bin/python docs/05_presentation/build_pptx.py
Out:  docs/05_presentation/llm-analytics-slides.pptx
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Emu, Inches, Pt

# --------------------------------------------------------------------------
# design tokens (mirrors the HTML deck)
# --------------------------------------------------------------------------

INK = RGBColor(0x0F, 0x17, 0x20)
INK2 = RGBColor(0x3B, 0x4A, 0x5A)
MUTED = RGBColor(0x6B, 0x7C, 0x8F)
LINE = RGBColor(0xD9, 0xE2, 0xEC)
PAPER = RGBColor(0xFF, 0xFF, 0xFF)
WASH = RGBColor(0xF4, 0xF7, 0xFA)
BRAND = RGBColor(0x0B, 0x6B, 0xCB)
BRAND_SOFT = RGBColor(0xE7, 0xF0, 0xFB)
OK = RGBColor(0x0F, 0x7B, 0x52)
OK_SOFT = RGBColor(0xE3, 0xF4, 0xEC)
WARN = RGBColor(0xB4, 0x69, 0x0E)
WARN_SOFT = RGBColor(0xFD, 0xF1, 0xDE)
STOP = RGBColor(0xB4, 0x23, 0x18)
STOP_SOFT = RGBColor(0xFD, 0xEC, 0xEB)
PHASE2 = RGBColor(0x6B, 0x4E, 0xA8)
PHASE2_SOFT = RGBColor(0xEF, 0xEA, 0xF9)

CODE_BG = RGBColor(0x0F, 0x17, 0x20)
CODE_FG = RGBColor(0xE6, 0xED, 0xF3)
CODE_COMMENT = RGBColor(0x7D, 0x8E, 0xA3)
CODE_KEY = RGBColor(0x8A, 0xB4, 0xF8)
CODE_STR = RGBColor(0xA5, 0xD6, 0xA7)

TH = "IBM Plex Sans Thai"
MONO = "Menlo"

W = Inches(13.333)
H = Inches(7.5)
M = Inches(0.62)          # outer margin
CONTENT_W = W - 2 * M

prs = Presentation()
prs.slide_width = W
prs.slide_height = H
BLANK = prs.slide_layouts[6]

_page = {"n": 0}


# --------------------------------------------------------------------------
# primitives
# --------------------------------------------------------------------------

def _style(run, *, size=12, bold=False, color=INK, font=TH):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def textbox(slide, x, y, w, h, *, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.paragraphs[0].alignment = align
    return box, tf


def para(tf, text="", *, first=False, size=12, bold=False, color=INK,
         font=TH, spacing=1.35, space_after=4, align=PP_ALIGN.LEFT):
    """Add a paragraph. `text` may be a plain string or a list of runs
    given as (text, {style overrides})."""
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.line_spacing = spacing
    p.space_after = Pt(space_after)

    chunks = text if isinstance(text, list) else [(text, {})]
    for chunk_text, override in chunks:
        run = p.add_run()
        run.text = chunk_text
        opts = {"size": size, "bold": bold, "color": color, "font": font}
        opts.update(override)
        _style(run, **opts)
    return p


def rect(slide, x, y, w, h, *, fill=PAPER, border=LINE, border_w=1.0,
         radius=0.06, shadow=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    shape.adjustments[0] = radius
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    if border is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = border
        shape.line.width = Pt(border_w)
    shape.shadow.inherit = shadow
    tf = shape.text_frame
    tf.word_wrap = True
    # safety net: if a Thai string runs longer than expected, PowerPoint
    # shrinks it rather than spilling outside the card.
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    tf.margin_left = tf.margin_right = Inches(0.16)
    tf.margin_top = tf.margin_bottom = Inches(0.12)
    tf.vertical_anchor = MSO_ANCHOR.TOP
    return shape


def pill(slide, x, y, text, *, fg=BRAND, bg=BRAND_SOFT, size=9.5,
         w=None, h=Inches(0.26)):
    """Small rounded 'tag' chip."""
    width = w or Inches(0.13 * len(text) + 0.28)
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, width, h)
    shape.adjustments[0] = 0.5
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg
    shape.line.fill.background()
    shape.shadow.inherit = False
    tf = shape.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = Inches(0.07)
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para(tf, [(text, {})], first=True, size=size, bold=True, color=fg,
         spacing=1.0, space_after=0, align=PP_ALIGN.CENTER)
    return shape


def card(slide, x, y, w, h, title, body, *, fill=PAPER, border=LINE,
         title_color=INK, title_size=13, body_size=10.5, body_color=INK2):
    shape = rect(slide, x, y, w, h, fill=fill, border=border)
    tf = shape.text_frame
    if title:
        para(tf, title, first=True, size=title_size, bold=True,
             color=title_color, spacing=1.25, space_after=5)
        para(tf, body, size=body_size, color=body_color, spacing=1.4)
    else:
        para(tf, body, first=True, size=body_size, color=body_color,
             spacing=1.4)
    return shape


def code_block(slide, x, y, w, h, lines):
    """lines: list of (text, kind) where kind in {'', 'c', 'k', 's'}."""
    shape = rect(slide, x, y, w, h, fill=CODE_BG, border=None, radius=0.045)
    tf = shape.text_frame
    tf.margin_left = tf.margin_right = Inches(0.18)
    tf.margin_top = tf.margin_bottom = Inches(0.14)
    palette = {"": CODE_FG, "c": CODE_COMMENT, "k": CODE_KEY, "s": CODE_STR}
    for i, item in enumerate(lines):
        if isinstance(item, str):
            item = [(item, "")]
        runs = [(t, {"color": palette[k]}) for t, k in item]
        para(tf, runs or [("", {})], first=(i == 0), size=9.5, font=MONO,
             color=CODE_FG, spacing=1.28, space_after=0)
    return shape


def table(slide, x, y, w, h, headers, rows, col_ratios, *,
          header_size=9.5, body_size=10, row_h=Inches(0.34)):
    shape = slide.shapes.add_table(len(rows) + 1, len(headers), x, y, w, h)
    tbl = shape.table
    tbl.first_row = True
    tbl.horz_banding = False

    total = sum(col_ratios)
    for i, ratio in enumerate(col_ratios):
        tbl.columns[i].width = Emu(int(w * ratio / total))

    tbl.rows[0].height = Inches(0.36)
    for r in range(1, len(rows) + 1):
        tbl.rows[r].height = row_h

    for c, head in enumerate(headers):
        cell = tbl.cell(0, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = WASH
        cell.margin_left = cell.margin_right = Inches(0.11)
        cell.margin_top = cell.margin_bottom = Inches(0.05)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        para(cell.text_frame, head, first=True, size=header_size, bold=True,
             color=MUTED, spacing=1.15, space_after=0)

    for r, row in enumerate(rows, start=1):
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = PAPER
            cell.margin_left = cell.margin_right = Inches(0.11)
            cell.margin_top = cell.margin_bottom = Inches(0.06)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            chunks = val if isinstance(val, list) else [(val, {})]
            para(cell.text_frame, chunks, first=True, size=body_size,
                 color=INK2, spacing=1.3, space_after=0)
    return tbl


# --------------------------------------------------------------------------
# slide chrome
# --------------------------------------------------------------------------

def new_slide(kicker=None, heading=None, sub=None):
    _page["n"] += 1
    slide = prs.slides.add_slide(BLANK)

    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = PAPER

    y = M
    if kicker:
        _, tf = textbox(slide, M, y, CONTENT_W, Inches(0.24))
        para(tf, kicker.upper(), first=True, size=10, bold=True, color=BRAND,
             spacing=1.0, space_after=0)
        y += Inches(0.3)
    if heading:
        _, tf = textbox(slide, M, y, CONTENT_W, Inches(0.55))
        para(tf, heading, first=True, size=25, bold=True, color=INK,
             spacing=1.15, space_after=0)
        y += Inches(0.58)
    if sub:
        _, tf = textbox(slide, M, y, Inches(10.6), Inches(0.44))
        para(tf, sub, first=True, size=11.5, color=MUTED, spacing=1.45,
             space_after=0)
        y += Inches(0.5)
    return slide, y + Inches(0.12)


def footer(slide, left, right=None):
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, M, H - Inches(0.72), CONTENT_W, Pt(1))
    line.fill.solid()
    line.fill.fore_color.rgb = LINE
    line.line.fill.background()
    line.shadow.inherit = False

    _, tf = textbox(slide, M, H - Inches(0.6), Inches(9.5), Inches(0.3))
    para(tf, left, first=True, size=9.5, color=MUTED, spacing=1.2,
         space_after=0)

    _, tf = textbox(slide, W - M - Inches(2.4), H - Inches(0.6), Inches(2.4),
                    Inches(0.3), align=PP_ALIGN.RIGHT)
    para(tf, right or f"{_page['n']} / 11", first=True, size=9.5, color=MUTED,
         spacing=1.2, space_after=0, align=PP_ALIGN.RIGHT)


# --------------------------------------------------------------------------
# slide 1 — title
# --------------------------------------------------------------------------

slide = prs.slides.add_slide(BLANK)
_page["n"] = 1
bg = slide.background.fill
bg.solid()
bg.fore_color.rgb = PAPER

accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.16), H)
accent.fill.solid()
accent.fill.fore_color.rgb = BRAND
accent.line.fill.background()
accent.shadow.inherit = False

_, tf = textbox(slide, Inches(0.95), Inches(1.5), Inches(11.2), Inches(0.3))
para(tf, "FINAL PROJECT · LAUNDRYTWIN", first=True, size=11,
     bold=True, color=BRAND, spacing=1.0, space_after=0)

_, tf = textbox(slide, Inches(0.95), Inches(1.95), Inches(11.2), Inches(1.7))
para(tf, "Safe AI Executive Assistant", first=True, size=40, bold=True,
     color=INK, spacing=1.12, space_after=2)
para(tf, "ที่ตอบจาก “ข้อมูลจริง” ไม่ใช่การเดา", size=40, bold=True, color=INK,
     spacing=1.12, space_after=0)

_, tf = textbox(slide, Inches(0.95), Inches(3.75), Inches(8.6), Inches(1.1))
para(tf, "ออกแบบโครงสร้างข้อมูล (Data Structure for LLM) และ Function Calling ที่ปลอดภัย "
         "เพื่อให้ผู้บริหารถามเป็นภาษาคนแล้วได้คำตอบที่ตรวจสอบแหล่งที่มาและขอบเขตได้ "
         "ภายใต้สิทธิ์ของผู้ใช้",
     first=True, size=13, color=MUTED, spacing=1.6, space_after=0)

meta = [
    ("ขอบเขต", "R08 · F-11 · F-07 · US-05 (MVP)"),
    ("ต่อยอด", "R09 · US-06 (Phase 2; F-12 คือ Weather Context)"),
    ("สถานะปัจจุบัน", "โครงสร้างข้อมูลและ allow-list บางส่วน implement แล้ว · local code/test evidence (ยืนยัน 2026-09-25)"),
    ("ฐานที่มีแล้ว", "Telemetry · Dashboard · RBAC · LINE Alert"),
]
mx = Inches(0.95)
for label, value in meta:
    _, tf = textbox(slide, mx, Inches(5.35), Inches(2.85), Inches(0.7))
    para(tf, label, first=True, size=9.5, bold=True, color=INK, spacing=1.2,
         space_after=3)
    para(tf, value, size=10.5, color=MUTED, spacing=1.35, space_after=0)
    mx += Inches(2.95)

footer(slide, "LaundryTwin — Smart Laundry Management and Analytics Platform",
       "1 / 11")

# --------------------------------------------------------------------------
# slide 2 — problem
# --------------------------------------------------------------------------

slide, y = new_slide(
    "ปัญหา",
    "ถ้าปล่อยให้ LLM ตอบเองแบบกล่องดำ จะพังตรงไหน",
    "โจทย์ไม่ใช่ “ต่อ LLM ได้ไหม” แต่คือ “จะเชื่อคำตอบได้อย่างไร”")

col_w = Inches(5.75)
gap_x = M + col_w + Inches(0.6)
box_h = Inches(3.5)

bad = rect(slide, M, y, col_w, box_h, fill=STOP_SOFT,
           border=RGBColor(0xF3, 0xBD, 0xB8))
tf = bad.text_frame
tf.margin_top = Inches(0.42)
para(tf, "Blackbox", first=True, size=15, bold=True, color=STOP, spacing=1.2,
     space_after=8)
para(tf, "โยนคำถามและข้อมูลดิบ (หรือสิทธิ์ query ฐานข้อมูล) ให้ LLM แล้วหวังว่าจะตอบถูก",
     size=11, color=INK2, spacing=1.5, space_after=10)
for item in [
    "ตัวเลขหลอน (hallucination) ตรวจย้อนกลับไม่ได้",
    "LLM เขียน SQL เอง = ช่องโหว่ข้ามสาขา ข้ามสิทธิ์",
    "ข้อมูลดิบยาวเกิน context และเปลืองต้นทุน",
    "ตอบไม่เหมือนเดิมทุกครั้ง (non-deterministic)",
    "ไม่มีหลักฐานว่าใครถามอะไร ระบบตอบจากอะไร",
]:
    para(tf, "•  " + item, size=11, color=INK2, spacing=1.5, space_after=3)
pill(slide, M + Inches(0.16), y + Inches(0.14), "แบบที่ไม่เอา", fg=STOP,
     bg=RGBColor(0xFA, 0xD9, 0xD6))

good = rect(slide, gap_x, y, col_w, box_h, fill=OK_SOFT,
            border=RGBColor(0xB7, 0xE0, 0xCD))
tf = good.text_frame
tf.margin_top = Inches(0.42)
para(tf, "Structured Function Calling", first=True, size=15, bold=True,
     color=OK, spacing=1.2, space_after=8)
para(tf, "LLM ทำหน้าที่แปลภาษาคนเป็นการเลือกเครื่องมือ และเรียบเรียงคำตอบเท่านั้น",
     size=11, color=INK2, spacing=1.5, space_after=10)
for item in [
    "ตัวเลขทั้งหมดมาจาก Analytics Function ที่ผ่านการทดสอบ",
    "Backend ตรวจ RBAC และ branch_id ก่อน query เสมอ",
    "ส่งเฉพาะข้อมูลสรุปแล้วขนาดเล็กเข้า context",
    "ผลลัพธ์คงที่ ทดสอบซ้ำได้ด้วย unit test",
     "Audit ปัจจุบันครอบคลุม grant / alert / settings; complete AI prompt-tool audit ยังเป็น target",
]:
    para(tf, "•  " + item, size=11, color=INK2, spacing=1.5, space_after=3)
pill(slide, gap_x + Inches(0.16), y + Inches(0.14), "แบบที่จะทำ", fg=OK,
     bg=RGBColor(0xC9, 0xEA, 0xDA))

arrow = slide.shapes.add_shape(
    MSO_SHAPE.RIGHT_ARROW, M + col_w + Inches(0.13),
    y + box_h / 2 - Inches(0.16), Inches(0.34), Inches(0.32))
arrow.fill.solid()
arrow.fill.fore_color.rgb = MUTED
arrow.line.fill.background()
arrow.shadow.inherit = False

footer(slide, "อ้างอิงข้อกำหนด R08 — “LLM must not execute arbitrary SQL”")

# --------------------------------------------------------------------------
# slide 3 — system overview
# --------------------------------------------------------------------------

slide, y = new_slide(
    "ภาพรวมระบบ",
    "ข้อมูลเดินทางจากเครื่องซักผ้า ถึงคำตอบผู้บริหารอย่างไร",
     "ภาพรวมนี้แยก current code ออกจาก target architecture: current มี scope, allow-list, direct reports และ envelope จริง; rollup กับ audit เต็มเป็นขั้นต่อไป")

row1 = [
    ("1. Edge / IoT", "Target physical path: Modbus/MQTT → telemetry event; current direct MQTT ingestion is descoped and IRIS is the ingestion boundary", False),
    ("2. Ingestion", "Target semantic model: validate branch/machine and derive MACHINE_CYCLE; current ETL normalizes usage, temperature, and weather into ClickHouse", False),
    ("3. Rollup", "Target: pre-aggregate by hour × branch × machine; current code still queries ClickHouse fact tables directly", True),
    ("4. Analytics Functions", "Target: semantic layer with a standard envelope; current has six MCP tools and `{ meta, data }`", True),
]
cw = Inches(2.95)
cx = M
for title, body, is_new in row1:
    shape = card(slide, cx, y, cw, Inches(1.72), title, body)
    if is_new:
        pill(slide, cx + cw - Inches(0.72), y + Inches(0.1), "ใหม่",
             w=Inches(0.56))
    cx += cw + Inches(0.13)

y2 = y + Inches(1.9)
row2 = [
    ("5. Assistant Orchestrator", "Current: รับคำถาม ตรวจ session ให้ LLM เลือก tool ตรวจ argument เรียก MCP แล้วเรียบเรียงคำตอบจากผลลัพธ์", True),
    ("6. Audit (บางส่วน)", "Current audit ครอบคลุม grant / alert / settings; target คือ prompt + tool + sanitized args + scope + result_ref + outcome", False),
    ("7. หน้าใช้งาน", "Current web มี branch/date, source/freshness/availability, analytics, alerts, AI settings/history, owner Playground และ legal navigation; browser QA pending", False),
]
cw3 = Inches(4.0)
cx = M
for title, body, is_new in row2:
    shape = card(slide, cx, y2, cw3, Inches(1.72), title, body)
    if is_new:
        pill(slide, cx + cw3 - Inches(0.72), y2 + Inches(0.1), "ใหม่",
             w=Inches(0.56))
    cx += cw3 + Inches(0.13)

footer(slide, "ทุกขั้นตอนบังคับขอบเขตสาขาที่ฝั่ง server เสมอ ตาม Data Contract")

# --------------------------------------------------------------------------
# slide 4 — activity flow
# --------------------------------------------------------------------------

slide, y = new_slide(
    "หัวใจของงาน · Activity Diagram 4",
    "ลำดับการทำงานเมื่อผู้ใช้ถามคำถามเชิงวิเคราะห์")


def lane(slide, x, y, w, h, label):
    shape = rect(slide, x, y, w, h, fill=WASH, border=LINE)
    _, tf = textbox(slide, x + Inches(0.18), y + Inches(0.13),
                    Inches(2.4), Inches(0.22))
    para(tf, label.upper(), first=True, size=8.5, bold=True, color=MUTED,
         spacing=1.0, space_after=0)
    return shape


def step(slide, x, y, w, h, tag, body, kind=""):
    fills = {"": (PAPER, LINE), "gate": (WARN_SOFT, RGBColor(0xF0, 0xC9, 0xA8)),
             "deny": (STOP_SOFT, RGBColor(0xF3, 0xBD, 0xB8)),
             "done": (OK_SOFT, RGBColor(0xB7, 0xE0, 0xCD))}
    fill, border = fills[kind]
    shape = rect(slide, x, y, w, h, fill=fill, border=border, radius=0.08)
    tf = shape.text_frame
    tf.margin_left = tf.margin_right = Inches(0.11)
    tf.margin_top = Inches(0.09)
    tf.margin_bottom = Inches(0.07)
    tag_color = {"": BRAND, "gate": WARN, "deny": STOP, "done": OK}[kind]
    para(tf, tag, first=True, size=8.5, bold=True, color=tag_color,
         spacing=1.1, space_after=3)
    para(tf, body, size=9, color=INK2, spacing=1.32, space_after=0)
    return shape


lane_h = Inches(1.15)
lane(slide, M, y, CONTENT_W, lane_h, "ผู้ใช้")
sx = M + Inches(0.18)
sw = Inches(5.85)
step(slide, sx, y + Inches(0.36), sw, Inches(0.68), "เริ่ม",
     "ถามเป็นภาษาคน เช่น “เดือนนี้รายได้สาขาสีลม เทียบกับเดือนที่แล้วเป็นอย่างไร”")
step(slide, sx + sw + Inches(0.25), y + Inches(0.36), sw, Inches(0.68), "จบ",
     "ได้คำตอบพร้อม ช่วงเวลา · ตัวชี้วัด · ข้อจำกัดของข้อมูล", "done")

y2 = y + lane_h + Inches(0.12)
lane(slide, M, y2, CONTENT_W, Inches(1.38), "Assistant Orchestrator")
steps2 = [
    ("ตรวจ 1", "session และสิทธิ์สมาชิกสาขา ถูกต้องหรือไม่", "gate"),
    ("บันทึก", "เก็บคำขอที่ sanitize แล้ว", ""),
    ("ตีความ", "จับ intent แล้วเสนอ tool จาก allow-list", ""),
    ("ตรวจ 2", "tool อยู่ใน allow-list หรือไม่", "gate"),
    ("สร้าง args", "แปลงเป็น argument แบบมีโครงสร้าง", ""),
    ("ตรวจ 3", "args ผ่าน schema validation แบบเข้มหรือไม่", "gate"),
    ("ตรวจ 4", "server ตรวจ tenant/branch scope; LINE scope ถูก sign จาก session", "gate"),
]
sw2 = Inches(1.63)
sx = M + Inches(0.18)
for tag, body, kind in steps2:
    step(slide, sx, y2 + Inches(0.4), sw2, Inches(0.84), tag, body, kind)
    sx += sw2 + Inches(0.09)

y3 = y2 + Inches(1.5)
lane(slide, M, y3, CONTENT_W, Inches(1.2), "Analytics Service")
steps3 = [
    ("ประมวลผล", "รันฟังก์ชันสถิติแบบ parameterized ไม่มี SQL จาก LLM", ""),
    ("บันทึก (target)", "เก็บชื่อ tool, sanitized args, scope, result_ref และ outcome; full AI audit ยังไม่มี", ""),
    ("ตรวจ 5", "ข้อมูลพอและตรวจย้อนกลับได้หรือไม่", "gate"),
]
sw3 = Inches(3.88)
sx = M + Inches(0.18)
for tag, body, kind in steps3:
    step(slide, sx, y3 + Inches(0.38), sw3, Inches(0.72), tag, body, kind)
    sx += sw3 + Inches(0.11)

y4 = y3 + Inches(1.32)
step(slide, M, y4, Inches(6.0), Inches(0.7), "ผ่านทั้งหมด",
     "เรียบเรียงคำตอบจากผลลัพธ์ของ tool เท่านั้น; audit ที่มีจริงยังไม่ครบ full AI call", "done")
step(slide, M + Inches(6.15), y4, Inches(5.95), Inches(0.7), "ไม่ผ่าน",
     "ปฏิเสธ / ไม่รองรับคำถาม / ข้อมูลไม่พอ พร้อมบันทึกเหตุผล", "deny")

footer(slide, "กล่องสีส้มคือประตูตรวจ ถ้าไม่ผ่านจะจบด้วยการปฏิเสธและบันทึก audit "
              "· ตรงกับ docs/02_architecture/data-and-activity-diagrams.md")

# --------------------------------------------------------------------------
# slide 5 — metric catalog
# --------------------------------------------------------------------------

slide, y = new_slide(
    "Data Structure for LLM · ชั้นที่ 1",
    "Metric Catalog — นิยามตัวชี้วัดให้ตรงกันก่อนเขียนโค้ด",
    "ถ้าคำว่า “รายได้” หรือ “อัตราการใช้งาน” ยังนิยามไม่ตรงกัน คำตอบของ AI จะเถียงกันเองตลอด")

rows = [
    ("revenue", "ผลรวมมูลค่ารอบซักที่มีหลักฐานการชำระเงิน", "สตางค์ (จำนวนเต็ม)",
     "เก็บเป็น satang ตามโค้ดเดิม กันปัญหาทศนิยม"),
    ("cycle_count", "จำนวนรอบซักที่จบสมบูรณ์", "รอบ",
     "นับจาก MACHINE_CYCLE ที่มีสถานะจบ"),
    ("utilization", "เวลาที่เครื่องทำงาน หารด้วยเวลาที่เปิดให้บริการ", "ร้อยละ",
     "ต้องระบุเวลาทำการของสาขาให้ชัด"),
    ("avg_cycle_minutes", "ค่าเฉลี่ยระยะเวลาต่อรอบ", "นาที",
     "ใช้ดูความผิดปกติของเครื่อง"),
    ("peak_hour", "ช่วงชั่วโมงที่มีรอบซักหนาแน่นสูงสุด", "ชั่วโมง 0–23",
     "ฐานของข้อเสนอโปรโมชัน (Phase 2)"),
    ("coinbox_estimate", "ประมาณการเหรียญสะสมในกล่อง", "สตางค์",
     "รีเซ็ตจาก event coinbox_open ที่แมปไว้ชัดเจนเท่านั้น ห้ามเดาจาก door_status"),
]
body = []
for name, definition, unit, note in rows:
    body.append([
        [(name, {"font": MONO, "size": 9.5, "color": INK})],
        definition,
        unit,
        note,
    ])
table(slide, M, y, CONTENT_W, Inches(2.9),
      ["ตัวชี้วัด", "นิยาม", "หน่วย", "หมายเหตุ"], body, [22, 30, 16, 32],
      row_h=Inches(0.42))

note_y = y + Inches(3.1)
card(slide, M, note_y, CONTENT_W, Inches(0.82), None,
     "กำหนดร่วมกันทุกตัวชี้วัด: เขตเวลา Asia/Bangkok · ขอบเขตช่วงเวลาแบบ [start, end) · "
     "ความละเอียด hour | day | month · และทุก query ต้องผูกกับ branch_id เสมอ",
     fill=BRAND_SOFT, border=RGBColor(0xC4, 0xDC, 0xF6), body_size=11.5)

footer(slide, "สอดคล้องกับ docs/03_data_contracts/data_contracts.md")

# --------------------------------------------------------------------------
# slide 6 — rollup
# --------------------------------------------------------------------------

slide, y = new_slide(
    "Data Structure for LLM · ชั้นที่ 2",
    "Pre-aggregated Rollup — target architecture ไม่ใช่ตารางที่ current code มีแล้ว")

left_w = Inches(5.6)
_, tf = textbox(slide, M, y, left_w, Inches(0.9))
para(tf, "แนวคิด target คือสรุปข้อมูลดิบระดับ event ให้เล็กและเร็วขึ้น "
         "จึงสรุปเป็นตารางระดับ ชั่วโมง × สาขา × เครื่อง ไว้ก่อน "
         "แต่ current implementation ยัง query ClickHouse fact tables โดยตรง",
     first=True, size=11.5, color=MUTED, spacing=1.55, space_after=0)

bullets = [
    ("เร็วและคงที่", "คำถามเดิมได้ตัวเลขเดิมทุกครั้ง ทดสอบด้วย unit test ได้"),
    ("เล็ก", "ส่งเข้า LLM เพียงหลักสิบแถว แทนข้อมูลดิบหลักแสน event"),
    ("ปลอดภัย", "ชั้นสรุปไม่มีข้อมูลส่วนบุคคล ลดความเสี่ยงข้อมูลรั่ว"),
    ("ต่อยอดง่าย", "รวมจากรายชั่วโมงเป็นรายวันหรือรายเดือนได้ทันที"),
]
by = y + Inches(1.0)
for head, body_text in bullets:
    card(slide, M, by, left_w, Inches(0.68), None,
         [(head + " — ", {"bold": True, "color": INK, "size": 11}),
          (body_text, {"size": 11})])
    by += Inches(0.74)

card(slide, M, by + Inches(0.06), left_w, Inches(0.82), None,
     "แนบตัวชี้วัดคุณภาพข้อมูลไปกับทุกแถว เช่น สัดส่วนชั่วโมงที่ telemetry ขาดหาย "
     "เพื่อให้คำตอบบอกข้อจำกัดได้ตรงความจริง",
     fill=WARN_SOFT, border=RGBColor(0xF0, 0xC9, 0xA8), body_size=11)

code_x = M + left_w + Inches(0.5)
code_w = CONTENT_W - left_w - Inches(0.5)
code_block(slide, code_x, y, code_w, Inches(4.9), [
    [("-- ตารางสรุปรายชั่วโมง (แนวคิด)", "c")],
    [("TABLE", "k"), (" metric_rollup_hourly (", "")],
    [("  branch_id             ", ""), ("-- บังคับ ใช้คุมขอบเขตสิทธิ์", "c")],
    "  machine_id",
    [("  bucket_start          ", ""), ("-- ต้นชั่วโมง Asia/Bangkok", "c")],
    "  cycle_count",
    "  revenue_satang",
    "  runtime_seconds",
    [("  open_seconds          ", ""), ("-- ใช้หาร utilization", "c")],
    "  telemetry_gap_seconds",
    [("  PRIMARY KEY", "k"), (" (branch_id, machine_id, bucket_start)", "")],
    ")",
    "",
    [("-- ตัวอย่างข้อมูลที่ LLM จะได้เห็น", "c")],
    [("{ ", ""), ('"branch_id"', "s"), (": ", ""), ('"BKK-SILOM"', "s"), (",", "")],
    [("  ", ""), ('"bucket_start"', "s"), (": ", ""),
     ('"2026-07-01T18:00+07:00"', "s"), (",", "")],
    [("  ", ""), ('"cycle_count"', "s"), (": 14,", "")],
    [("  ", ""), ('"revenue_satang"', "s"), (": 84000,", "")],
    [("  ", ""), ('"utilization_pct"', "s"), (": 62.5 }", "")],
])

footer(slide, "ชั้นนี้คือภาษากลางระหว่างฐานข้อมูลกับ AI")

# --------------------------------------------------------------------------
# slide 7 — allow-listed tools
# --------------------------------------------------------------------------

slide, y = new_slide(
    "Data Structure for LLM · ชั้นที่ 3",
    "Current MCP Tools — LLM เลือกได้เฉพาะ 6 เครื่องมือที่ server allow-list ไว้",
    "LLM ไม่ได้รับสิทธิ์เขียน SQL แต่ได้รับเมนูฟังก์ชันที่ผ่านการทดสอบแล้วเท่านั้น")

tool_rows = [
    ("get_revenue_daily", "รายได้และรอบรายวัน; ต้องมี revenue scope"),
    ("get_cycles_daily", "รอบรายวันและ duration เฉลี่ย"),
    ("get_utilization_heatmap", "ช่วงเวลา/เครื่องสำหรับ utilization"),
    ("get_temperature_curve", "อุณหภูมิตามช่วงเวลา ไม่เกิน 5,000 จุด"),
    ("get_weather_usage_correlation", "correlation เท่านั้น ไม่ใช่ forecast"),
    ("get_off_peak_windows", "baseline ช่วงว่าง Phase 2"),
]
left_w = Inches(6.0)
table(slide, M, y, left_w, Inches(2.7), ["Current tool", "ใช้ตอบคำถามแบบไหน"],
      [[[(n, {"font": MONO, "size": 9.5, "color": INK})], d]
       for n, d in tool_rows],
      [42, 58], row_h=Inches(0.38))

card(slide, M, y + Inches(2.85), left_w, Inches(0.72), None,
     "คำถามที่ไม่มี tool รองรับ ระบบตอบว่า “ยังไม่รองรับ” ซึ่งดีกว่าเดาแล้วผิด",
     fill=WASH, border=LINE, body_size=11)

code_x = M + left_w + Inches(0.5)
code_w = CONTENT_W - left_w - Inches(0.5)
code_block(slide, code_x, y, code_w, Inches(2.42), [
    [("// นิยาม argument แบบเข้มงวด (zod / JSON Schema)", "c")],
    "get_cycles_daily({",
    [("  from:     ", ""), ('"2026-08-01"', "s")],
    [("  to:       ", ""), ('"2026-08-31"', "s")],
    [("  branchId: ", ""), ('"b1"', "s"), (" // server-bound scope; ไม่มี accessScope argument", "c")],
    "})",
])
code_block(slide, code_x, y + Inches(2.56), code_w, Inches(1.95), [
    [("// ลำดับการตรวจฝั่ง backend (ไม่เชื่อ LLM)", "c")],
    [("1.", "k"), (" parse ด้วย schema — ผิดรูปคือปฏิเสธ", "")],
    [("2.", "k"), (" server-bound scope = ขอบเขตจริงจาก session/grants", "")],
    [("3.", "k"), (" ถ้า scope ไม่ผ่าน = ปฏิเสธก่อน query; full AI audit ยังเป็น target", "")],
    [("4.", "k"), (" จำกัดความยาวช่วงเวลาและจำนวนแถว", "")],
    [("5.", "k"), (" query แบบ parameterized เท่านั้น", "")],
])

card(slide, M, y + Inches(3.35), CONTENT_W, Inches(0.72), None,
     [("กติกาเหล็ก: ", {"bold": True, "color": STOP, "size": 11.5}),
      ("ต่อให้ LLM ขอสาขาที่ไม่มีสิทธิ์ backend จะตัดออกเสมอ "
       "Manager จึงดึงข้อมูลข้ามสาขาไม่ได้ ตามตัวอย่างใน R08", {"size": 11.5})],
     fill=STOP_SOFT, border=RGBColor(0xF3, 0xBD, 0xB8))

footer(slide, "ตรงกับ F-11 และ US-05")

# --------------------------------------------------------------------------
# slide 8 — result envelope
# --------------------------------------------------------------------------

slide, y = new_slide(
    "Data Structure for LLM · ชั้นที่ 4",
    "Result Envelope — current `{ meta, data }` และ target ที่จะเพิ่ม trace fields")

code_w = Inches(6.15)
code_block(slide, M, y, code_w, Inches(4.55), [
    "{",
    [('  "meta"', "s"), (": {", "")],
    [('    "range"', "s"), (": { ", ""), ('"from"', "s"), (": ", ""), ('"2026-08-01"', "s"), (", ", ""), ('"to"', "s"), (": ", ""), ('"2026-08-31"', "s"), (" },", "")],
    [('    "branchId"', "s"), (": ", ""), ('"b1"', "s"), (",", "")],
    [('    "dataSource"', "s"), (": ", ""), ('"real"', "s")],
    "  },",
    [('  "data"', "s"), (": [ ... ]", "")],
    "}",
])

rx = M + code_w + Inches(0.45)
rw = CONTENT_W - code_w - Inches(0.45)
notes = [
    ("Current envelope",
     "meta มี range, branchId, dataSource และบาง tool มี method/rules/caveats; data เป็นแถวผลลัพธ์",
     PAPER, LINE, INK),
    ("Target envelope ที่ยังไม่ ship",
     "จะเพิ่ม scope, period, metric, unit, coverage, source, generated_at และ result_ref เพื่อให้ตรวจย้อนกลับได้",
     PAPER, LINE, INK),
    ("ทำไมยังไม่เรียกว่า audit เต็ม",
     "Current schema มี audit สำหรับ grants, alerts, settings; ยังไม่มี prompt/tool-call/result audit",
     PAPER, LINE, INK),
    ("Target system prompt",
     "“ตอบจากผลลัพธ์ของ tool เท่านั้น ห้ามประมาณเอง ถ้าข้อมูลไม่พอให้บอกว่าไม่พอ”",
     OK_SOFT, RGBColor(0xB7, 0xE0, 0xCD), OK),
]
ny = y
for title, body_text, fill, border, tcolor in notes:
    card(slide, rx, ny, rw, Inches(1.06), title, body_text, fill=fill,
         border=border, title_color=tcolor, title_size=12, body_size=10.5)
    ny += Inches(1.16)

footer(slide, "Current `{ meta, data }` ใช้กับทุก tool; richer envelope เป็น target")

# --------------------------------------------------------------------------
# slide 9 — security
# --------------------------------------------------------------------------

slide, y = new_slide(
    "ความปลอดภัยและการตรวจสอบ",
    "ความปลอดภัยที่มีแล้ว และสิ่งที่ยังเป็น target")

pillars = [
    ("เสา 1", "RBAC + Branch Scope",
     "สิทธิ์ถูกตัดสินจาก session ฝั่ง server ไม่ใช่จากสิ่งที่ LLM ขอ ทุก query ผูกกับ branch_id ที่ตัดกับสิทธิ์จริงแล้ว"),
    ("เสา 2", "Audit (บางส่วน)",
     "Current audit entries ครอบคลุม grants, alerts, settings; complete AI prompt/tool-call/result audit ยังเป็น target"),
    ("เสา 3", "Golden-question Eval (target)",
     "ชุดคำถามมาตรฐานพร้อมคำตอบที่รู้ค่าจริง จะใช้จับ regression และการหลอนตัวเลข; ยังไม่ใช่ current automated evidence"),
]
cw = Inches(3.96)
cx = M
for tag, title, body_text in pillars:
    card(slide, cx, y, cw, Inches(1.68), title, body_text)
    pill(slide, cx + cw - Inches(0.78), y + Inches(0.1), tag, w=Inches(0.62))
    cx += cw + Inches(0.13)

ty = y + Inches(1.88)
sec_rows = [
    ("Manager ถามข้อมูลสาขาที่ไม่ได้ดูแล",
     "ปฏิเสธก่อน query; local test ยืนยัน scope gate แต่ยังไม่ใช่ full AI audit"),
    ("ถามเรื่องที่ยังไม่มี tool รองรับ",
     "ตอบว่ายังไม่รองรับ ไม่พยายามเดาคำตอบ"),
    ("ช่วงเวลาที่ถามมี telemetry ขาดหายมาก",
     "Current อาจมี freshness/caveats; target จะเพิ่ม coverage และบอกชัดว่าข้อมูลไม่พอ"),
    ("ถามคำถามเดิมซ้ำสองครั้ง",
     "ได้ตัวเลขเดิม เพราะตัวเลขมาจากฟังก์ชัน ไม่ใช่จากการสร้างข้อความ"),
]
table(slide, M, ty, CONTENT_W, Inches(2.0),
      ["สถานการณ์ทดสอบ", "ผลลัพธ์ที่ระบบต้องทำ"], sec_rows, [38, 62],
      row_h=Inches(0.4))

footer(slide, "ชุดหลักฐานสำหรับการสาธิตหน้าชั้น")

# --------------------------------------------------------------------------
# slide 10 — plan
# --------------------------------------------------------------------------

slide, y = new_slide(
    "แผนดำเนินงาน",
    "ลำดับงานและสิ่งส่งมอบ",
    "Target roadmap: current มี direct reports, RBAC และ 6 MCP tools; semantic rollup, richer envelope, full AI audit และ golden eval ยังต้องทำ")

plan_rows = [
    ("1", "สรุป Metric Catalog และเขตเวลา หน่วย", "เอกสารนิยามตัวชี้วัดใน docs/03_data_contracts/", "ทีมเห็นตรงกันทุกนิยาม"),
    ("2", "สร้างชั้น Rollup รายชั่วโมง (target)", "ตารางสรุป งานสรุปข้อมูล และ unit test", "ยังไม่มีใน current schema; ต้องพิสูจน์กับ fact tables"),
    ("3", "เขียน Analytics Functions (target expansion)", "เก็บ current 6 tools; target อาจเพิ่ม semantic functions", "current tool allow-list มี test; target ต้องมี test ครบ"),
    ("4", "บังคับ RBAC และ branch scope", "ชั้นตรวจสิทธิ์ก่อนทุก query", "เคส Manager ข้ามสาขาต้องถูกปฏิเสธ"),
    ("5", "ต่อ LLM ด้วย tool schema", "Orchestrator, allow-list และ system prompt", "ไม่มีเส้นทางใดที่ LLM ส่ง SQL ได้"),
    ("6", "AI audit เต็มและหน้าใช้งาน (target)", "เก็บ prompt/tool/result ครบ + UI ถาม-ตอบ", "current มี audit เฉพาะ grant/alert/settings; full audit ยังไม่มี"),
    ("7", "ชุดทดสอบ Golden question (target)", "ชุดคำถามและคำตอบมาตรฐาน", "ยังไม่ใช่ current automated evidence"),
]
table(slide, M, y, CONTENT_W, Inches(3.1),
      ["ลำดับ", "งาน", "สิ่งส่งมอบ", "ตรวจรับด้วย"],
      [[[(n, {"bold": True, "color": BRAND})], w, d, a]
       for n, w, d, a in plan_rows],
      [8, 30, 34, 28], row_h=Inches(0.36))

sy = y + Inches(3.3)
card(slide, M, sy, Inches(5.95), Inches(0.95), "Current MVP ที่มี local evidence",
     "R08 · F-11 · US-05 — direct reports, RBAC scope, 6 MCP tools และ LINE bot; production/LINE/browser E2E ยัง pending",
     fill=BRAND_SOFT, border=RGBColor(0xC4, 0xDC, 0xF6),
     title_color=BRAND)
shape = card(slide, M + Inches(6.15), sy, Inches(5.95), Inches(0.95),
             "Phase 2 / target",
             "R09 off-peak baseline ใช้ percentile; weather เป็น correlation ไม่ใช่ forecast; richer semantic layer และ full audit ยังต้องพัฒนา",
             fill=PHASE2_SOFT, border=RGBColor(0xD8, 0xCD, 0xF0),
             title_color=PHASE2)
pill(slide, M + Inches(6.15) + Inches(5.95) - Inches(1.35), sy + Inches(0.12),
     "ยังไม่ทำตอนนี้", fg=PHASE2, bg=RGBColor(0xE2, 0xD9, 0xF6), w=Inches(1.2))

footer(slide, "ทุกแถวผูกกับรหัสข้อกำหนดใน docs/04_traceability/RTM_matrix.md")

# --------------------------------------------------------------------------
# slide 11 — summary
# --------------------------------------------------------------------------

slide, y = new_slide("สรุป", "ข้อความที่อยากให้จำ")

points = [
    ("1", "LLM ไม่แตะข้อมูลดิบ",
     "ทำหน้าที่แปลคำถามและเรียบเรียงคำตอบ ตัวเลขทั้งหมดมาจากฟังก์ชันที่ทดสอบแล้ว"),
    ("2", "สิทธิ์ตัดสินที่ server",
     "ขอบเขตสาขาถูกบังคับก่อน query เสมอ; LINE bot ส่ง scope ที่ server derive และ sign ไว้"),
    ("3", "ตรวจย้อนกลับได้ในระดับฐาน",
     "มีช่วงเวลา source freshness และ caveat; full AI result_ref audit เป็น target"),
]
cw = Inches(3.96)
cx = M
for num, title, body_text in points:
    shape = rect(slide, cx, y, cw, Inches(2.35), fill=OK_SOFT,
                 border=RGBColor(0xB7, 0xE0, 0xCD))
    tf = shape.text_frame
    tf.margin_top = Inches(0.22)
    para(tf, num, first=True, size=34, bold=True, color=OK, spacing=1.0,
         space_after=6)
    para(tf, title, size=14, bold=True, color=INK, spacing=1.25, space_after=6)
    para(tf, body_text, size=11, color=INK2, spacing=1.45, space_after=0)
    cx += cw + Inches(0.13)

_, tf = textbox(slide, M, y + Inches(2.75), Inches(10.4), Inches(1.3))
para(tf, [("เป้าหมายของหัวข้อนี้ไม่ใช่การใส่ AI ให้ระบบดูทันสมัย "
           "แต่คือการออกแบบโครงสร้างข้อมูลที่ทำให้คำตอบของ AI ", {}),
          ("เชื่อถือได้ ตรวจสอบได้ และปลอดภัยตามสิทธิ์ผู้ใช้", {"bold": True, "color": INK}),
          (" ซึ่งเป็นสิ่งที่ระบบธุรกิจจริงต้องการ", {})],
     first=True, size=15, color=INK2, spacing=1.6, space_after=0)

footer(slide, "LaundryTwin — Smart Laundry Management and Analytics Platform",
       "จบการนำเสนอ")

# --------------------------------------------------------------------------

out = Path(__file__).resolve().parent / "llm-analytics-slides.pptx"
prs.core_properties.title = "LaundryTwin — Safe AI Executive Assistant (R08 / F-11)"
prs.core_properties.author = "LaundryTwin final project"
prs.save(out)
print(f"wrote {out}  ({len(prs.slides.__iter__.__self__._sldIdLst)} slides)")
